"""Host-side Word exporter; Pandoc preserves mathematics as editable Word equations."""
import argparse
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills" / "universal-exam-cram-coach"))
from coach.guide import digest, save
from coach.guide_chapters import validate


def style_document(path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(path)
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.3)
        sec.left_margin = sec.right_margin = Cm(2.4)
    for name in ("Normal", "Body Text", "First Paragraph"):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(12)
        st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(8)
    for name, size in (("Title", 22), ("Subtitle", 12), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
    for math_run in doc.element.iter(qn("m:r")):
        props = math_run.find(qn("w:rPr"))
        if props is None:
            props = OxmlElement("w:rPr")
            math_run.insert(0, props)
        fonts = OxmlElement("w:rFonts")
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + key), "Cambria Math")
        props.append(fonts)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "27" if any(p.tag == qn("m:oMathPara") for p in math_run.iterancestors()) else "24")
        props.append(size)
        # LibreOffice renders the OMML divides glyph as a slash; use a literal conditional bar.
        for node in math_run.iter(qn("m:t")):
            if node.text:
                node.text = node.text.replace("∣", "|")
    if doc.paragraphs and doc.paragraphs[0].style.name == "Heading 1":
        doc.paragraphs[0].style = doc.styles["Title"]
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 2":
            paragraph.paragraph_format.space_before = Pt(20)
            paragraph.paragraph_format.keep_with_next = True
        runs = [run for run in paragraph.runs if run.text.strip()]
        if runs and all(run.bold for run in runs):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph._p.find(qn("m:oMathPara")) is not None:
            previous = paragraph._p.getprevious()
            if previous is not None and previous.tag == qn("w:p"):
                from docx.text.paragraph import Paragraph
                Paragraph(previous, paragraph._parent).paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(12)
    for table in doc.tables:
        table.autofit = True
        props = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + edge)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                el.set(qn("w:" + key), value)
            borders.append(el)
        props.append(borders)
        for cell in table.rows[0].cells:
            sh = OxmlElement("w:shd")
            sh.set(qn("w:fill"), "E9EEF3")
            cell._tc.get_or_add_tcPr().append(sh)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = 1
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.35
                    paragraph.paragraph_format.space_after = Pt(6)
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = 1
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        p._p.append(field)
    doc.save(path)


def export(run, chapter, draft, review, output, preview=False):
    draft, out, run = Path(draft).resolve(), Path(output).resolve(), Path(run).resolve()
    if out.suffix.lower() != ".docx" or out.exists():
        raise ValueError("choose a new .docx output path")
    if not preview:
        if not review:
            raise ValueError("formal export requires content review")
        report = validate(run, chapter, draft, review)
        if not report["valid"]:
            raise ValueError("chapter validation failed: " + "; ".join(report["errors"]))
    converter = shutil.which("pandoc")
    if not converter:
        raise RuntimeError("Pandoc is required for editable Word equations")
    text = draft.read_text(encoding="utf-8")
    if preview and "样稿" not in text[:500]:
        raise ValueError("preview must visibly identify itself as 样稿 in the opening")
    # Only local images are permitted; exporting never fetches remote course assets.
    images = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", text)
    for value in images:
        image = value.strip().strip("<>")
        if "://" in image or not (draft.parent / image).is_file():
            raise ValueError("image must be an existing local file: " + image)
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="guide-export-", dir=out.parent) as tmp:
        temporary = Path(tmp) / "chapter.docx"
        subprocess.run([converter, str(draft), "--from=markdown", "--to=docx", "--standalone",
                        "--resource-path=" + str(draft.parent), "--output=" + str(temporary)], check=True)
        style_document(temporary)
        if not preview:
            report = validate(run, chapter, draft, review)
            if not report["valid"]:
                raise ValueError("sources or content changed during export")
        temporary.replace(out)
    receipt = {"status": "preview" if preview else "content_validated",
               "draft_sha256": digest(draft), "docx_sha256": digest(out),
               "sources_digest": digest(run / "sources.json"), "chapter": chapter,
               "visual_review": "pending", "output": str(out)}
    save(out.with_suffix(".receipt.json"), receipt)
    return receipt


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--run", required=True)
    p.add_argument("--chapter", required=True)
    p.add_argument("--draft", required=True)
    p.add_argument("--review")
    p.add_argument("--out", required=True)
    p.add_argument("--preview", action="store_true", help="explicitly labelled sample; never a validated chapter")
    a = p.parse_args()
    try:
        print(json.dumps(export(a.run, a.chapter, a.draft, a.review, a.out, a.preview), ensure_ascii=False, indent=2))
        return 0
    except (ValueError, OSError, RuntimeError, subprocess.CalledProcessError) as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
